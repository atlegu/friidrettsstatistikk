"""Avviksliste: basen mot de godkjente norske mastersrekordene (Excel).

Leser rapporten fra masters_rekorder_fra_liste.py (hvilke rekorder som ble
funnet eller lagt inn) og sammenligner hver godkjent hovedrekord med det
rekordsiden viser (masters_rekorder() per kjønn, klasse og ute/inne).

Arkene:
  Oversikt                    tellinger
  Bedre i basen               basen har et bedre resultat enn rekorden, med
                              sannsynlig årsak satt maskinelt og lenke til utøveren
  Godkjent rekord vises ikke  rekorden står ikke på rekordsiden
  Annen dato i basen          samme resultat, annen dato
  Navn og fødselsår           samme person, annet navn eller fødselsår i basen

Bare lesing fra basen; skriver én Excel-fil.

    python masters_avviksliste.py opprydding/masters_rekorder_<tid>.csv
"""

import argparse
import csv
import logging
import re
import sys
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import masters_rekorder_fra_liste as m

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
log = logging.getLogger(__name__)
sb = m.sb
HER = Path(__file__).parent
LISTE = {'mvu': 'Menn ute', 'kvu': 'Kvinner ute', 'mvi': 'Menn inne', 'kvi': 'Kvinner inne'}
URL = 'https://www.friidrettsresultater.no/utover/'
LISTEDATO = {'u': '2025-12-31', 'i': '2026-06-30'}
ALIAS_NAVN = {'Halldis Marit (Waaga) (Nagell-)Dahl': 'Halldis Marit Nagell-Dahl',
              'Astrid Førde': 'Astrid Mæhre Førde Johannessen',
              'Trond Erik Nilsen': 'Trond-Erik Nilsen', 'Lisbet Pedersen': 'Lisbeth Pedersen'}
# Navn i rekordlista som Atle har bekreftet er samme person som i basen
# (17.10.2026). Rune Bråten Liholt er født 1976; «(77)» i innelista er feil.
BEKREFTET_SAMME = {
    ('Aksel Røste', '35'), ('Anna Magrethe Tråve', '55'), ('Arthur Lillefosse', '44'), ('Astrid Førde', '33'),
    ('Bjørn Morten Moss', '60'), ('Claire Phytian', '73'), ('Dag Kåshagen', '60'), ('Elin Margrete Dølør', '65'),
    ('Elisabeth Wormsen', '55'), ('Elise Marie Wåle', '42'), ('Emma Grønmo', '44'), ('Eva E. Nordqvist', '63'),
    ('Grete Rivenes', '48'), ('Halldis Marit (Waaga) (Nagell-)Dahl', '42'), ('Inger Elisabeth Holter Holmgrunn', '49'),
    ('Inger-Lise Sundheim', '55'), ('Janice Terjesen Flaathe', '59'), ('Katarina Sederholm Hoff', '68'),
    ('Kirsten Marathon Melkevik', '70'), ('Kirsti Mathilde Johnsen', '53'), ('Kjell Adamski', '50'),
    ('Kjell Steinset', '36'), ('Klaus Arthur Hansen', '58'), ('Lisbet Pedersen', '75'),
    ('May Britt Hasselquist Borge', '61'), ('Maylem Lauvstad', '73'), ('Oddbjørg Haakensveen', '45'),
    ('Ragnhild Hanstvedt', '34'), ('Reidar Dahl', '48'), ('Rune Bråten Liholt', '77'), ('Sigrun R. Sangvik', '57'),
    ('Sofie A. Jakobsen', '43'), ('Sverre Kristoffer Jensen', '66'), ('Synøve Brox', '59'),
    ('Toril Lauritsen Nyborg', '67'), ('Trond Erik Nilsen', '72'), ('Åse Nyland', '33'),
}

MERKNAD = {'Erik Tysse': 'Basen har 11:28.45 samme dag (Framolekene 2015). En av tidene er feil. Godkjent tid er ikke lagt inn.',
           'Erling Andersen': 'Lista: M40-44 i 1990, men oppgitt født 1960 (30 år). Feil klasse, dato eller fødselsår i lista.'}


def liste(fil):
    return LISTE[fil[15:18]]


def ord_i(navn):
    return set(re.sub(r'[()\-.]', ' ', (navn or '').lower()).split())


def aarsak(x):
    b = x['basen']
    if b['date'] > LISTEDATO[x['fil'][17]]:
        return '1 Nyere enn lista – mulig ny rekord'
    p, tid, ov = b['performance'], b['result_type'] == 'time', x['overskrift']
    dist = re.search(r'(\d+)\s?(m|km)\b', ov)
    if tid and ('.' in x['perf'] or ',' in x['perf']) and (
            re.match(r'^\d+:\d\d(:\d\d)?$', p) or re.match(r'^\d{3,}\.\d+$', p)
            or (re.match(r'^\d\.\d\d$', p) and dist and int(dist.group(1)) >= 200)):
        return '2 Avkortet, avrundet eller feil tidsformat i basen'
    if tid and re.search(r':00$', p):
        return '2 Avkortet, avrundet eller feil tidsformat i basen'
    if tid and re.match(r'^\d+\.\d$', p) and dist and int(dist.group(1)) < 800:
        return '3 Håndtid i basen som ikke er merket'
    off, v = int(x['pv']), b['performance_value']
    if (off / v > 1.08) if tid else (v / off > 1.12):
        return '4 Stor forskjell (over 8 % tid / 12 % lengde) – kontroller særlig nøye'
    if ord_i(b['athlete_name']) & ord_i(x['navn']):
        return '5 Samme utøver – bedre resultat i basen enn det godkjente'
    return '6 Annen utøver – ikke godkjent? Kontroller fødselsår og statsborgerskap'


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('rapport')
    ap.add_argument('--ut', default=str(HER.parent / 'NFIFProsjekt' / f'Mastersrekorder_avvik_{date.today()}.xlsx'))
    args = ap.parse_args()
    rap = list(csv.DictReader(open(args.rapport, encoding='utf-8'), delimiter=';'))
    events = {e['code']: e for e in sb.table('events').select('code,result_type').execute().data}

    # Det rekordsiden viser
    vaare = {}
    for kj in 'MF':
        for inne in (False, True):
            for k in range(35, 110, 5):
                for x in sb.rpc('masters_rekorder', {'p_kjonn': kj, 'p_klasse': f'V{k}', 'p_inne': inne}).execute().data:
                    vaare[(kj, inne, k, x['event_code'])] = x
    rader = []
    for x in rap:
        if not (x['hoved'] == 'True' and x['koder'] and x['pv']):
            continue
        kj, inne = ('M' if x['fil'][15] == 'm' else 'F'), x['fil'][17] == 'i'
        k = int(re.match(r'[MK](\d+)', x['klasse']).group(1))
        koder = x['koder'].split(',')
        tid = events[koder[0]]['result_type'] == 'time'
        off, beste = int(x['pv']), None
        for c in koder:
            v = vaare.get((kj, inne, k, c))
            if v and (beste is None or ((v['performance_value'] < beste['performance_value']) if tid
                                        else (v['performance_value'] > beste['performance_value']))):
                beste = v
        if not beste:
            s = 'mangler'
        elif beste['performance_value'] == off:
            s = 'lik'
        else:
            s = 'bedre' if (beste['performance_value'] < off) == tid else 'dårligere'
        rader.append({**x, 's': s, 'basen': beste})
    log.info(Counter(r['s'] for r in rader))

    bedre = [x for x in rader if x['s'] == 'bedre']
    info = {}
    for i in {x['basen']['athlete_id'] for x in bedre}:
        f = sb.table('results').select('date').eq('athlete_id', i).order('date').limit(1).execute().data
        a = sb.table('athletes').select('nationality').eq('id', i).execute().data
        info[i] = {'forste': f[0]['date'] if f else None, 'nasjon': a[0]['nationality'] if a else None}
    for x in bedre:
        x['aarsak'] = aarsak(x)
    bedre.sort(key=lambda x: (x['aarsak'], x['fil'], x['overskrift'], x['klasse']))

    wb = Workbook()
    fet, fyll = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='192E5D')
    lenke = Font(color='0563C1', underline='single')

    def ark(navn, hoder, data, bredder, lenkekol=None):
        ws = wb.create_sheet(navn)
        ws.append(hoder)
        for c in ws[1]:
            c.font, c.fill, c.alignment = fet, fyll, Alignment(wrap_text=True, vertical='center')
        for rad, url in data:
            ws.append(rad)
            if lenkekol and url:
                c = ws.cell(row=ws.max_row, column=lenkekol)
                c.hyperlink, c.font = url, lenke
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(hoder, 1):
            ws.column_dimensions[get_column_letter(i)].width = bredder.get(h, 12)
        return len(data)

    n1 = ark('Bedre i basen',
             ['Sannsynlig årsak', 'Liste', 'Øvelse', 'Klasse', 'Godkjent rekord', 'Vind', 'Rekordholder', 'Født', 'Dato',
              'Bedre i basen', 'Vind (basen)', 'Utøver i basen', 'Født (basen)', 'Alder', 'Dato (basen)', 'Stevne (basen)',
              'Første resultat i basen', 'Nasjon (basen)'],
             [([x['aarsak'][2:], liste(x['fil']), x['overskrift'], x['klasse'], x['perf'], x['vind'], x['navn'], x['fodt'], x['dato'],
                x['basen']['performance'], x['basen']['wind'], x['basen']['athlete_name'], x['basen']['fodt'],
                int(x['basen']['date'][:4]) - x['basen']['fodt'] if x['basen']['fodt'] else None, x['basen']['date'],
                x['basen']['meet_name'], info[x['basen']['athlete_id']]['forste'], info[x['basen']['athlete_id']]['nasjon']],
               URL + x['basen']['athlete_id']) for x in bedre],
             {'Sannsynlig årsak': 50, 'Øvelse': 32, 'Rekordholder': 26, 'Utøver i basen': 28, 'Stevne (basen)': 44,
              'Godkjent rekord': 10, 'Bedre i basen': 10, 'Dato': 11, 'Dato (basen)': 11, 'Første resultat i basen': 12},
             lenkekol=12)

    ikke = [x for x in rader if x['s'] in ('dårligere', 'mangler')]
    n2 = ark('Godkjent rekord vises ikke',
             ['Liste', 'Øvelse', 'Klasse', 'Godkjent rekord', 'Rekordholder', 'Født', 'Dato', 'Basen viser', 'Utøver i basen', 'Merknad'],
             [([liste(x['fil']), x['overskrift'], x['klasse'], x['perf'], x['navn'], x['fodt'], x['dato'],
                x['basen']['performance'] if x['basen'] else '', x['basen']['athlete_name'] if x['basen'] else '',
                MERKNAD.get(x['navn'], '')], None) for x in ikke],
             {'Øvelse': 32, 'Rekordholder': 22, 'Utøver i basen': 22, 'Merknad': 80})

    dato = [r for r in rap if r['status'] == 'finnes med annen dato']
    n3 = ark('Annen dato i basen',
             ['Liste', 'Øvelse', 'Klasse', 'Resultat', 'Utøver', 'Født', 'Sted', 'Dato i lista', 'Dato i basen', 'Merknad'],
             [([liste(r['fil']), r['overskrift'], r['klasse'] or '(tillegg)', r['perf'], r['navn'], r['fodt'], r['sted'],
                r['dato'], r['merknad'],
                'Mesterskap over flere dager lagres ofte med startdatoen' if r['sted'] and ',' in r['sted'] else ''], None)
              for r in dato],
             {'Øvelse': 32, 'Utøver': 26, 'Sted': 18, 'Merknad': 50})

    # Samme person, annet navn eller fødselsår
    sett, data = set(), []
    for r in rap:
        koblet = r['koblet'] or ''
        if not (r['status'] == 'finnes (annen navneform i basen)' or koblet.startswith(('alias', 'navn, fødselsår', 'fornavn+etternavn'))):
            continue
        if (r['navn'], r['fodt']) in BEKREFTET_SAMME:
            continue
        if r['athlete_id']:
            a = sb.table('athletes').select('full_name,birth_year').eq('id', r['athlete_id']).execute().data
            i_basen = f"{a[0]['full_name']} ({a[0]['birth_year']})" if a else ''
            lik = a and a[0]['full_name'].lower() == r['navn'].lower() and str(a[0]['birth_year'])[-2:] == r['fodt'][-2:]
        else:
            i_basen, lik = r['merknad'] or ALIAS_NAVN.get(r['navn'], ''), False
        if lik or (r['navn'], r['fodt'], i_basen) in sett:
            continue
        sett.add((r['navn'], r['fodt'], i_basen))
        data.append(([liste(r['fil']), r['navn'], r['fodt'], i_basen,
                      'Fødselsår avviker' if 'fødselsår' in koblet else 'Navneform eller skrivemåte avviker',
                      r['overskrift'], r['perf'], r['dato']], None))
    n4 = ark('Navn og fødselsår', ['Liste', 'Navn i lista', 'Født (lista)', 'I basen (navn, født)', 'Avvik',
                                   'Eksempel: øvelse', 'Resultat', 'Dato'], data,
             {'Navn i lista': 34, 'I basen (navn, født)': 40, 'Avvik': 34, 'Eksempel: øvelse': 30})

    ws = wb.active
    ws.title = 'Oversikt'
    c = Counter(x['aarsak'] for x in bedre)
    linjer = [('Avvik mellom friidrettsresultater.no og de godkjente norske mastersrekordene', None),
              ('Rekordlister: Stein Fossen (kappgang Erling Andersen). Ute pr. 31.12.2025, inne pr. 30.06.2026. '
               f'Kontrollert {date.today():%d.%m.%Y}.', None), ('', None),
              ('Hovedrekorder kontrollert', len(rader)),
              ('Stemmer med basen (samme resultat vises på rekordsiden)', sum(1 for x in rader if x['s'] == 'lik')), ('', None),
              ('Basen har et bedre resultat enn den godkjente rekorden (ark «Bedre i basen»)', n1)]
    linjer += [('   ' + k[2:], v) for k, v in sorted(c.items())]
    linjer += [('Godkjent rekord vises ikke på rekordsiden (ark «Godkjent rekord vises ikke»)', n2),
               ('Samme resultat, men annen dato i basen (ark «Annen dato i basen»)', n3),
               ('Samme person, men annet navn eller fødselsår i basen, ikke bekreftet (ark «Navn og fødselsår»)', n4),
               (f'   Bekreftet samme person (17.10.2026), ikke tatt med', len(BEKREFTET_SAMME)), ('', None),
               ('Årsakene er satt maskinelt og er et utgangspunkt for vurdering, ikke en konklusjon.', None),
               ('«Stor forskjell»: over 8 % bedre tid eller 12 % lengre/høyere enn rekorden. Ofte feil fødselsår, '
                'øvelse eller utøver, men kan være ekte.', None),
               ('Utøvernavnene i arket «Bedre i basen» lenker til utøversiden.', None)]
    for a, b in linjer:
        ws.append([a, b])
    ws['A1'].font = Font(bold=True, size=14, color='192E5D')
    ws.column_dimensions['A'].width, ws.column_dimensions['B'].width = 95, 10
    wb.save(args.ut)
    log.info(f"Lagret {args.ut}")
    for a, b in linjer[3:-4]:
        if a:
            log.info(f"{a}: {b}")


if __name__ == '__main__':
    main()
